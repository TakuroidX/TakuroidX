"""Export the input sheets of the Excel MATH spec back to math JSON.

Workflow for tuning math in Excel:
  1. In docs/math_spec.xlsx, edit the blue cells in Reel Strips / Paytable / Base Parameters (RTP recalculates immediately)
  2. python3 tools/export_math.py docs/math_spec.xlsx math/neon_fortune.json
  3. ./tools/make_specs.sh  (SDK verification → regenerate Excel)

Usage: python3 tools/export_math.py <xlsx> <out.json>
"""
import json
import sys

from openpyxl import load_workbook

sys.path.insert(0, __import__("os").path.dirname(__file__))
from design_reels import compact_arrays  # noqa: E402


def rows(ws, first_row, ncols, first_col=1):
    r = first_row
    while True:
        vals = [ws.cell(row=r, column=first_col + c).value for c in range(ncols)]
        if vals[0] is None:
            return
        yield vals
        r += 1


def export(xlsx):
    wb = load_workbook(xlsx, data_only=False)
    p = {ws_row[0]: ws_row for ws_row in rows(wb["Base Parameters"], 4, 5)}

    def param(label):
        return p[label][1]

    symbols = [{"id": i, "name": n, "type": t} for i, n, t in rows(wb["Symbols"], 4, 3)]

    pay_ws = wb["Paytable"]
    paytable = {}
    for vals in rows(pay_ws, 5, 6):
        paytable[vals[0]] = [int(v or 0) for v in vals[1:]]
    # SCATTER row: the data row right after the header row that follows the normal-symbol table
    scatter_row = 5 + len(paytable) + 2
    scatter_pays = [int(pay_ws.cell(row=scatter_row, column=c).value or 0) for c in range(2, 7)]

    reel_ws = wb["Reel Strips"]
    n_reels = int(param("Reels"))
    reels = []
    for r in range(n_reels):
        strip = []
        row = 5
        while (v := reel_ws.cell(row=row, column=2 + r).value) is not None:
            strip.append(str(v).strip())
            row += 1
        reels.append(strip)

    paylines = [[int(x) for x in vals[1:]] for vals in rows(wb["Paylines"], 4, 1 + n_reels)]

    return {
        "id": "neon_fortune",
        "name": "Neon Fortune",
        "version": "1.0.0",
        "grid": {"reels": n_reels, "rows": int(param("Rows"))},
        "symbols": symbols,
        "wild": {"symbol": param("WILD symbol ID")},
        "scatter": {"symbol": param("SCATTER symbol ID"), "pays": scatter_pays, "triggerCount": int(param("FS trigger SCATTER count"))},
        "freeSpins": {
            "awardSpins": int(param("FS spins awarded")),
            "retriggerSpins": int(param("FS retrigger spins")),
            "multiplier": int(param("FS multiplier")),
        },
        "bet": {
            "lines": int(param("Lines")),
            "lineBetLevels": [int(x) for x in p["Line bet levels"][1:5] if x is not None],
            "defaultLevel": int(param("Default level (0-based)")),
        },
        "paytable": paytable,
        "paylines": paylines,
        "reels": reels,
    }


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    model = export(sys.argv[1])
    # Keep id/name/version from the existing JSON if it exists
    try:
        with open(sys.argv[2], encoding="utf-8") as f:
            old = json.load(f)
        for k in ("id", "name", "version"):
            model[k] = old.get(k, model[k])
    except FileNotFoundError:
        pass
    text = compact_arrays(json.dumps(model, indent=2, ensure_ascii=False))
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        f.write(text + "\n")
    print(f"wrote {sys.argv[2]}")


if __name__ == "__main__":
    main()
