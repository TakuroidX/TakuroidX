"""Common styles for the Excel specs."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT_NAME = "Meiryo UI"

NAVY = "1F2A44"
INPUT_BLUE = "0000FF"
LINK_GREEN = "008000"

TITLE = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
SUBTITLE = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
HEADER = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY = Font(name=FONT_NAME, size=10)
BOLD = Font(name=FONT_NAME, size=10, bold=True)
INPUT = Font(name=FONT_NAME, size=10, color=INPUT_BLUE)
LINK = Font(name=FONT_NAME, size=10, color=LINK_GREEN)
NOTE = Font(name=FONT_NAME, size=9, italic=True, color="666666")

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF9C4")
TOTAL_FILL = PatternFill("solid", fgColor="E3E8F2")
SECTION_FILL = PatternFill("solid", fgColor="D6DCE9")

THIN = Side(style="thin", color="B0B7C3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)

PCT4 = "0.0000%"
PCT2 = "0.00%"
PROB = "0.00000000"
NUM = "#,##0"
NUM2 = "#,##0.00"


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = NOTE
    ws.sheet_view.showGridLines = False
    # Printing: landscape, fit to 1 page wide
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def header(ws, row, col, labels, widths=None):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font = HEADER
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BOX
    if widths:
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(widths):
            if w:
                ws.column_dimensions[get_column_letter(col + i)].width = w


def cell(ws, row, col, value, font=BODY, fmt=None, fill=None, align=None, border=True):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = align or (LEFT if isinstance(value, str) and not str(value).startswith("=") else CENTER)
    if border:
        c.border = BOX
    return c


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w
