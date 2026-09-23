"""Generate the Excel MATH spec (docs/math_spec.xlsx).

Inputs:
  math/neon_fortune.json      … math definition (source of truth)
  math/neon_fortune.sim.json  … slotsim report output (omitted if absent)

Every theoretical value is an Excel formula. Edit the "Reel Strips" or "Paytable" sheet in Excel
and the RTP recalculates. To feed edits back into the SDK, use tools/export_math.py.

Usage: python3 tools/build_math_spec.py
"""
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

import xlstyle as S
from xlstyle import cell, header, title, widths

ROOT = Path(__file__).resolve().parent.parent
MATH = ROOT / "math" / "neon_fortune.json"
SIM = ROOT / "math" / "neon_fortune.sim.json"
OUT = ROOT / "docs" / "math_spec.xlsx"

MAX_STOPS = 100          # rows kept for the reel strip input area
P = "'Base Parameters'"  # sheet references are always quoted
SYM = "'Symbols'"
PAY = "'Paytable'"
REEL = "'Reel Strips'"
CNT = "'Symbol Counts'"
LINE = "'Line Calc'"
SCAT = "'Scatter Calc'"
FS = "'Free Spin Calc'"
RTP = "'RTP Breakdown'"
SIMS = "'Simulation Results'"

SYMBOL_NOTES = {
    "WILD": "Substitutes for every symbol except SCATTER. Reels 2–4 only. No pay of its own.",
    "SCATTER": "3+ anywhere pays x total bet + awards free spins.",
}


def main():
    m = json.loads(MATH.read_text(encoding="utf-8"))
    sim = json.loads(SIM.read_text(encoding="utf-8")) if SIM.exists() else None
    wb = Workbook()
    wb.remove(wb.active)

    normal = [s["id"] for s in m["symbols"] if s["type"] == "normal"]
    symbols = [s["id"] for s in m["symbols"]]
    n_reels = m["grid"]["reels"]

    cover(wb, m)
    wb.create_sheet("Summary")  # contents are filled last (they reference other sheets' row numbers)
    params(wb, m)
    symbol_sheet(wb, m)
    paytable(wb, m, normal)
    reels(wb, m)
    paylines(wb, m)
    counts(wb, symbols, n_reels)
    line_calc(wb, normal, n_reels, len(symbols))
    scatter_calc(wb, n_reels, len(symbols))
    fs_calc(wb, len(normal))
    rtp_breakdown(wb, len(normal))
    if sim:
        sim_sheet(wb, sim)
        verify(wb, sim)
    summary(wb["Summary"], m, sim)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula when opened
    wb.save(OUT)
    print(f"wrote {OUT}")


# ---------------------------------------------------------------- cover / summary
def cover(wb, m):
    ws = wb.create_sheet("Cover")
    title(ws, f"{m['name']}  MATH Spec")
    widths(ws, {"A": 22, "B": 60, "C": 18, "D": 40})
    rows = [
        ("Game ID", m["id"]), ("Game name", m["name"]), ("Math version", m["version"]),
        ("Created", date.today().isoformat()), ("Author", "TakuroidX (AI Director) / Claude Code"),
        ("Purpose", "Hobby / learning project. Not intended for real-money gambling."),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        cell(ws, i, 1, k, S.BOLD, fill=S.SECTION_FILL)
        cell(ws, i, 2, v)

    ws["A11"] = "Sheets"
    ws["A11"].font = S.SUBTITLE
    header(ws, 12, 1, ["Sheet", "Contents"])
    sheets = [
        ("Summary", "Key figures (RTP, hit frequency, FS frequency, volatility)"),
        ("Base Parameters", "Grid, lines, bet, free spin settings (input)"),
        ("Symbols", "Symbol list and rules"),
        ("Paytable", "Pays (input)"),
        ("Reel Strips", "Reel strip for each reel (input)"),
        ("Paylines", "Row pattern of the 20 lines (input)"),
        ("Symbol Counts", "Count of each symbol per reel (COUNTIF)"),
        ("Line Calc", "Exact line-pay probabilities and RTP"),
        ("Scatter Calc", "SCATTER count distribution (32-pattern enumeration) and FS trigger probability"),
        ("Free Spin Calc", "Expected FS spins with retriggers, and FS RTP"),
        ("RTP Breakdown", "RTP breakdown by component"),
        ("Simulation Results", "Full-cycle enumeration and Monte Carlo results from the SDK (slotsim)"),
        ("Verification", "Excel formulas vs SDK cross-check"),
    ]
    for i, (k, v) in enumerate(sheets, start=13):
        cell(ws, i, 1, k)
        cell(ws, i, 2, v)

    r = 13 + len(sheets) + 1
    ws.cell(row=r, column=1, value="Legend").font = S.SUBTITLE
    legend = [
        ("Blue text + yellow fill", "Input values (design parameters). Edit these cells", S.INPUT, S.INPUT_FILL),
        ("Black text", "Formulas within the sheet", S.BODY, None),
        ("Green text", "Formulas referencing another sheet", S.LINK, None),
        ("Gray italics", "Values pasted from SDK output (source noted)", S.NOTE, None),
    ]
    for i, (k, v, f, fill) in enumerate(legend, start=r + 1):
        cell(ws, i, 1, k, f, fill=fill)
        cell(ws, i, 2, v)

    r += len(legend) + 2
    ws.cell(row=r, column=1, value="Revision History").font = S.SUBTITLE
    header(ws, r + 1, 1, ["Version", "Changes", "Date", "Author"])
    for c, v in enumerate([m["version"], "Initial version (5x3 / 20 lines / FS x3)", date.today().isoformat(), "TakuroidX"], start=1):
        cell(ws, r + 2, c, v)


def summary(ws, m, sim):
    title(ws, "Summary", "Theoretical values are formulas linked to the calculation sheets. Measured values come from the Simulation Results sheet.")
    widths(ws, {"A": 34, "B": 18, "C": 18, "D": 50})
    header(ws, 4, 1, ["Item", "Theory (Excel)", "Measured (SDK)", "Notes"])
    rows = [
        ("Total RTP", f"={RTP}!B9", f"={SIMS}!B6" if sim else None, S.PCT4, "Target: see Base Parameters"),
        ("  Line pays", f"={RTP}!B5", None, S.PCT4, ""),
        ("  Scatter pays", f"={RTP}!B6", None, S.PCT4, ""),
        ("  Free spins", f"={RTP}!B8", f"={SIMS}!B9" if sim else None, S.PCT4, ""),
        ("FS trigger frequency (1 in N)", f"={SCAT}!B{SCAT_TRIG_ROW + 1}", f"=1/{SIMS}!B11" if sim else None, S.NUM2, "Probability of 3+ SCATTERs"),
        ("Expected FS spins (incl. retriggers)", f"={FS}!B9", f"={SIMS}!B12" if sim else None, "0.000", ""),
        ("Hit frequency (base game)", None, f"={SIMS}!B10" if sim else None, S.PCT2, "Probability of any win in a spin"),
        ("Volatility (per-round SD)", None, f"={SIMS}!B14" if sim else None, "0.000", "Standard deviation of the win multiple per round"),
        ("Max base game win (x total bet)", None, f"={SIMS}!E8" if sim else None, S.NUM2, "Exact value from full-cycle enumeration"),
        ("Max round win observed (x total bet)", None, f"={SIMS}!B15" if sim else None, S.NUM2, "Monte Carlo observation (not a theoretical max)"),
    ]
    for i, (k, th, me, fmt, note) in enumerate(rows, start=5):
        cell(ws, i, 1, k, S.BOLD if not k.startswith("  ") else S.BODY)
        cell(ws, i, 2, th, S.LINK, fmt)
        cell(ws, i, 3, me, S.LINK, fmt)
        cell(ws, i, 4, note)
    r = 5 + len(rows) + 1
    ws.cell(row=r, column=1, value="Game specification").font = S.SUBTITLE
    specs = [
        ("Layout", f"{m['grid']['reels']} reels x {m['grid']['rows']} rows, {m['bet']['lines']} fixed lines, left-to-right pays"),
        ("Bet", f"line bet {m['bet']['lineBetLevels']} credits x {m['bet']['lines']} lines"),
        ("WILD", SYMBOL_NOTES["WILD"]),
        ("SCATTER", "3/4/5 anywhere = " + "/".join(str(x) for x in m["scatter"]["pays"][2:]) + " x total bet"),
        ("Free spins", f"{m['scatter']['triggerCount']}+ SCATTER → {m['freeSpins']['awardSpins']} spins, all wins x{m['freeSpins']['multiplier']}, "
                  f"{m['scatter']['triggerCount']}+ during FS adds {m['freeSpins']['retriggerSpins']} spins (unlimited)"),
        ("Max win cap", "None (base game max is shown above)"),
    ]
    for i, (k, v) in enumerate(specs, start=r + 1):
        cell(ws, i, 1, k, S.BOLD, fill=S.SECTION_FILL)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        cell(ws, i, 2, v)


# ---------------------------------------------------------------- inputs
PARAM_ROWS = {}  # key -> row


def params(wb, m):
    ws = wb.create_sheet("Base Parameters")
    title(ws, "Base Parameters", "Blue cells are inputs. Every calculation sheet references these cells.")
    widths(ws, {"A": 28, "B": 14, "C": 10, "D": 10, "E": 10, "F": 44})
    header(ws, 3, 1, ["Item", "Value", "", "", "", "Notes"])
    rows = [
        ("reels", "Reels", m["grid"]["reels"], "Changing this requires rebuilding the calc sheets"),
        ("rows", "Rows", m["grid"]["rows"], "Visible symbols per reel"),
        ("lines", "Lines", m["bet"]["lines"], "Fixed lines. Must match the Paylines sheet"),
        ("levels", "Line bet levels", None, "Credits per line (4 levels)"),
        ("defaultLevel", "Default level (0-based)", m["bet"]["defaultLevel"], ""),
        ("award", "FS spins awarded", m["freeSpins"]["awardSpins"], "On a base game trigger"),
        ("retrigger", "FS retrigger spins", m["freeSpins"]["retriggerSpins"], "On 3+ during FS"),
        ("mult", "FS multiplier", m["freeSpins"]["multiplier"], "Applies to every FS win"),
        ("trigger", "FS trigger SCATTER count", m["scatter"]["triggerCount"], ""),
        ("target", "Target RTP", 0.96, "Design target"),
        ("wild", "WILD symbol ID", m["wild"]["symbol"], ""),
        ("scatter", "SCATTER symbol ID", m["scatter"]["symbol"], ""),
    ]
    for i, (key, label, value, note) in enumerate(rows, start=4):
        PARAM_ROWS[key] = i
        cell(ws, i, 1, label, S.BOLD)
        if key == "levels":
            for j, lv in enumerate(m["bet"]["lineBetLevels"]):
                cell(ws, i, 2 + j, lv, S.INPUT, fill=S.INPUT_FILL)
        else:
            cell(ws, i, 2, value, S.INPUT, S.PCT2 if key == "target" else None, S.INPUT_FILL)
        cell(ws, i, 6, note)


def pref(key):
    return f"{P}!$B${PARAM_ROWS[key]}"


def symbol_sheet(wb, m):
    ws = wb.create_sheet("Symbols")
    title(ws, "Symbols")
    widths(ws, {"A": 12, "B": 12, "C": 10, "D": 18, "E": 60})
    header(ws, 3, 1, ["ID", "Name", "Type", "Reels", "Notes"])
    for i, s in enumerate(m["symbols"], start=4):
        where = ", ".join(str(r + 1) for r, reel in enumerate(m["reels"]) if s["id"] in reel)
        cell(ws, i, 1, s["id"], S.INPUT, fill=S.INPUT_FILL)
        cell(ws, i, 2, s["name"], S.INPUT, fill=S.INPUT_FILL)
        cell(ws, i, 3, s["type"], S.INPUT, fill=S.INPUT_FILL)
        cell(ws, i, 4, where)
        cell(ws, i, 5, SYMBOL_NOTES.get(s["id"], "Normal symbol (left-to-right pays)"))


PAY_FIRST = 5          # first normal-symbol row in the Paytable sheet
SCAT_PAY_ROW = 0       # set in paytable()


def paytable(wb, m, normal):
    global SCAT_PAY_ROW
    ws = wb.create_sheet("Paytable")
    title(ws, "Paytable", "Normal symbols: x line bet (left-to-right matches from reel 1). SCATTER: x total bet (anywhere).")
    widths(ws, {"A": 14, "B": 10, "C": 10, "D": 10, "E": 10, "F": 10, "G": 40})
    header(ws, 4, 1, ["Symbol", "1 of a kind", "2 of a kind", "3 of a kind", "4 of a kind", "5 of a kind", "Notes"])
    for i, sym in enumerate(normal, start=PAY_FIRST):
        cell(ws, i, 1, sym, S.BOLD)
        for k, v in enumerate(m["paytable"][sym]):
            cell(ws, i, 2 + k, v, S.INPUT, S.NUM, S.INPUT_FILL)
        cell(ws, i, 7, "")
    r = PAY_FIRST + len(normal) + 1
    header(ws, r, 1, ["SCATTER", "1", "2", "3", "4", "5", "x total bet"])
    SCAT_PAY_ROW = r + 1
    cell(ws, SCAT_PAY_ROW, 1, m["scatter"]["symbol"], S.BOLD)
    for k, v in enumerate(m["scatter"]["pays"]):
        cell(ws, SCAT_PAY_ROW, 2 + k, v, S.INPUT, S.NUM, S.INPUT_FILL)
    cell(ws, SCAT_PAY_ROW, 7, "x total bet; triggers free spins")
    note = ws.cell(row=SCAT_PAY_ROW + 2, column=1,
                   value="Wins on the same line pay only the highest. Line pays and scatter pays are added together.")
    note.font = S.NOTE


REEL_FIRST = 5


def reels(wb, m):
    ws = wb.create_sheet("Reel Strips")
    title(ws, "Reel Strips", f"Stops 0 to (length-1). Row r shows strip[(stop+r) mod length]. Up to {MAX_STOPS} rows can be entered.")
    widths(ws, {"A": 10, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 4, "H": 50})
    cell(ws, 3, 1, "Length", S.BOLD, fill=S.TOTAL_FILL)
    for r in range(len(m["reels"])):
        col = 2 + r
        L = ws.cell(row=1, column=col).column_letter
        cell(ws, 3, col, f"=COUNTA({L}{REEL_FIRST}:{L}{REEL_FIRST + MAX_STOPS - 1})", S.BOLD, fill=S.TOTAL_FILL)
    header(ws, 4, 1, ["Stop"] + [f"Reel {i + 1}" for i in range(len(m["reels"]))])
    longest = max(len(r) for r in m["reels"])
    for s in range(longest):
        cell(ws, REEL_FIRST + s, 1, s)
        for r, reel in enumerate(m["reels"]):
            if s < len(reel):
                cell(ws, REEL_FIRST + s, 2 + r, reel[s], S.INPUT, fill=S.INPUT_FILL)
    rules = [
        "Design rules (checked by the SDK's MathModel.Validate):",
        "- Don't put WILD on reel 1 (premise of the exact calculation)",
        "- Keep SCATTERs on the same reel at least 3 apart (at most 1 visible per reel)",
        "- The same symbol is never adjacent (a design choice in design_reels.py)",
    ]
    for i, t in enumerate(rules):
        ws.cell(row=4 + i, column=8, value=t).font = S.NOTE if i else S.BOLD
    # Color-code the symbols
    colors = {"WILD": "E1BEE7", "SCATTER": "FFE082", "SEVEN": "FFCDD2", "DIAMOND": "B3E5FC"}
    rng = f"B{REEL_FIRST}:F{REEL_FIRST + MAX_STOPS - 1}"
    for sym, color in colors.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{sym}"'], fill=PatternFill("solid", fgColor=color)))
    ws.freeze_panes = f"B{REEL_FIRST}"


def paylines(wb, m):
    ws = wb.create_sheet("Paylines")
    title(ws, "Paylines", "Row numbers: 0 = top, 1 = middle, 2 = bottom. Shape is a formula (▔ = top, ─ = middle, ▁ = bottom).")
    widths(ws, {"A": 8, "B": 8, "C": 8, "D": 8, "E": 8, "F": 8, "G": 16})
    header(ws, 3, 1, ["Line", "Reel 1", "Reel 2", "Reel 3", "Reel 4", "Reel 5", "Shape"])
    for i, line in enumerate(m["paylines"], start=4):
        cell(ws, i, 1, i - 3)
        for r, row in enumerate(line):
            cell(ws, i, 2 + r, row, S.INPUT, fill=S.INPUT_FILL)
        cell(ws, i, 7, "=" + "&".join(f'CHOOSE({c}{i}+1,"▔","─","▁")' for c in "BCDEF"))


# ---------------------------------------------------------------- calculations
CNT_FIRST = 5
CNT_LEN_ROW = 0


def counts(wb, symbols, n_reels):
    global CNT_LEN_ROW
    ws = wb.create_sheet("Symbol Counts")
    title(ws, "Symbol Counts per Reel", "Aggregated from the Reel Strips sheet with COUNTIF.")
    widths(ws, {"A": 14, "B": 10, "C": 10, "D": 10, "E": 10, "F": 10, "G": 10})
    header(ws, 4, 1, ["Symbol"] + [f"Reel {i + 1}" for i in range(n_reels)] + ["Total"])
    last = CNT_FIRST + len(symbols) - 1
    for i, sym in enumerate(symbols):
        row = CNT_FIRST + i
        cell(ws, row, 1, f"={SYM}!A{4 + i}", S.LINK)
        for r in range(n_reels):
            L = "BCDEF"[r]
            cell(ws, row, 2 + r, f"=COUNTIF({REEL}!{L}${REEL_FIRST}:{L}${REEL_FIRST + MAX_STOPS - 1},$A{row})", S.LINK)
        cell(ws, row, 2 + n_reels, f"=SUM(B{row}:F{row})")
    CNT_LEN_ROW = last + 1
    cell(ws, CNT_LEN_ROW, 1, "Length", S.BOLD, fill=S.TOTAL_FILL)
    for r in range(n_reels + 1):
        L = "BCDEFG"[r]
        cell(ws, CNT_LEN_ROW, 2 + r, f"=SUM({L}{CNT_FIRST}:{L}{last})", S.BOLD, fill=S.TOTAL_FILL)
    cell(ws, CNT_LEN_ROW + 1, 1, "Check", S.BOLD)
    for r in range(n_reels):
        L = "BCDEF"[r]
        cell(ws, CNT_LEN_ROW + 1, 2 + r, f'=IF({L}{CNT_LEN_ROW}={REEL}!{L}3,"OK","Unknown symbol")', S.LINK)
    ws.cell(row=CNT_LEN_ROW + 2, column=1,
            value="'OK' when every entry in the reel strip matches a symbol ID in the Symbols sheet.").font = S.NOTE


LINE_FIRST = 7
LINE_TOTAL_ROW = 0


def line_calc(wb, normal, n_reels, n_symbols):
    global LINE_TOTAL_ROW
    ws = wb.create_sheet("Line Calc")
    title(ws, "Line Calc (exact)",
          "Reel 1 has no WILD, so each line's paying symbol = reel 1's symbol. P(k+ in a row) = n1/L1 × Π(ni+Wi)/Li. "
          "Reels are independent and stops are equally likely, so every line has the same distribution → line RTP = per-line expectation (x line bet) = total-bet ratio.")
    cnt_rng = f"{CNT}!$A${CNT_FIRST}:$A${CNT_FIRST + n_symbols - 1}"

    cell(ws, 3, 1, "Length", S.BOLD, fill=S.TOTAL_FILL)
    cell(ws, 4, 1, "WILD count", S.BOLD, fill=S.TOTAL_FILL)
    for r in range(n_reels):
        L = "BCDEF"[r]
        cell(ws, 3, 2 + r, f"={CNT}!{L}{CNT_LEN_ROW}", S.LINK, fill=S.TOTAL_FILL)
        cell(ws, 4, 2 + r, f"=INDEX({CNT}!{L}${CNT_FIRST}:{L}${CNT_FIRST + n_symbols - 1},MATCH({pref('wild')},{cnt_rng},0))",
             S.LINK, fill=S.TOTAL_FILL)

    groups = [
        ("Effective count (reel 1 = symbol / reels 2+ = symbol + WILD)", 2),
        ("P(k+ in a row from the left)", 7),
        ("P(exactly k in a row)", 12),
        ("Pay (x line bet)", 17),
        ("Contribution = probability × pay", 22),
    ]
    for label, col in groups:
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 4)
        c = ws.cell(row=5, column=col, value=label)
        c.font = S.HEADER
        c.fill = S.HEADER_FILL
        c.alignment = S.CENTER
    labels = ["Symbol"]
    labels += [f"R{i + 1}" for i in range(5)]
    labels += [f"≥{k}" for k in range(1, 6)]
    labels += [f"={k}" for k in range(1, 6)]
    labels += [f"{k} kind" for k in range(1, 6)]
    labels += [f"{k} kind" for k in range(1, 6)]
    labels += ["Symbol total", "Hit rate / line"]
    header(ws, 6, 1, labels)
    ws.column_dimensions["A"].width = 12
    for col in range(2, 29):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 11

    pay_rng = f"{PAY}!$A${PAY_FIRST}:$A${PAY_FIRST + len(normal) - 1}"
    for i in range(len(normal)):
        row = LINE_FIRST + i
        cell(ws, row, 1, f"={PAY}!A{PAY_FIRST + i}", S.LINK)
        # Effective count
        for r in range(n_reels):
            L = "BCDEF"[r]
            base = f"INDEX({CNT}!{L}${CNT_FIRST}:{L}${CNT_FIRST + n_symbols - 1},MATCH($A{row},{cnt_rng},0))"
            cell(ws, row, 2 + r, f"={base}" if r == 0 else f"={base}+{L}$4", S.LINK)
        # P(>=k): G..K
        cell(ws, row, 7, f"=B{row}/B$3", fmt=S.PROB)
        for k in range(1, 5):
            prev = "GHIJK"[k - 1]
            cnt = "BCDEF"[k]
            cell(ws, row, 7 + k, f"={prev}{row}*{cnt}{row}/{cnt}$3", fmt=S.PROB)
        # P(=k): L..P
        for k in range(5):
            ge = "GHIJK"[k]
            if k < 4:
                cell(ws, row, 12 + k, f"={ge}{row}-{'GHIJK'[k + 1]}{row}", fmt=S.PROB)
            else:
                cell(ws, row, 12 + k, f"={ge}{row}", fmt=S.PROB)
        # Pay: Q..U
        for k in range(5):
            pl = "BCDEF"[k]
            cell(ws, row, 17 + k, f"=INDEX({PAY}!{pl}${PAY_FIRST}:{pl}${PAY_FIRST + len(normal) - 1},MATCH($A{row},{pay_rng},0))",
                 S.LINK, S.NUM)
        # Contribution: V..Z
        for k in range(5):
            pc = "LMNOP"[k]
            pp = "QRSTU"[k]
            cell(ws, row, 22 + k, f"={pc}{row}*{pp}{row}", fmt=S.PROB)
        cell(ws, row, 27, f"=SUM(V{row}:Z{row})", S.BOLD, S.PCT4)
        cell(ws, row, 28, f"=SUMPRODUCT(L{row}:P{row}*(Q{row}:U{row}>0))", fmt=S.PCT4)

    LINE_TOTAL_ROW = LINE_FIRST + len(normal)
    t = LINE_TOTAL_ROW
    cell(ws, t, 1, "Total", S.BOLD, fill=S.TOTAL_FILL)
    for col in range(2, 27):
        cell(ws, t, col, None, fill=S.TOTAL_FILL)
    cell(ws, t, 27, f"=SUM(AA{LINE_FIRST}:AA{t - 1})", S.BOLD, S.PCT4, S.TOTAL_FILL)
    cell(ws, t, 28, f"=SUM(AB{LINE_FIRST}:AB{t - 1})", S.BOLD, S.PCT4, S.TOTAL_FILL)
    ws.cell(row=t + 2, column=1, value="Line RTP (total-bet ratio)").font = S.BOLD
    cell(ws, t + 2, 5, f"=AA{t}", S.BOLD, S.PCT4)
    ws.cell(row=t + 3, column=1,
            value="Expected pay per line (x line bet) × number of lines ÷ total bet (line bet × number of lines) = expected pay per line").font = S.NOTE
    ws.freeze_panes = "B7"


SCAT_ENUM_FIRST = 13
SCAT_DIST_FIRST = 0
SCAT_TOTAL_ROW = 0
SCAT_TRIG_ROW = 0


def scatter_calc(wb, n_reels, n_symbols):
    global SCAT_DIST_FIRST, SCAT_TOTAL_ROW, SCAT_TRIG_ROW
    ws = wb.create_sheet("Scatter Calc")
    title(ws, "Scatter Calc (exact)",
          "SCATTERs are at least 3 apart, so at most 1 is visible per reel. P(visible on reel i) = rows × count / length. "
          "Enumerate all 2^5 = 32 visible/not-visible combinations to get the distribution of the SCATTER count.")
    widths(ws, {"A": 16, "B": 14, "C": 12, "D": 14, "E": 12, "F": 12, "G": 10, "H": 16})
    header(ws, 4, 1, ["Reel", "SCATTER count", "Length", "Visible prob p"])
    cnt_rng = f"{CNT}!$A${CNT_FIRST}:$A${CNT_FIRST + n_symbols - 1}"
    for r in range(n_reels):
        row = 5 + r
        cell(ws, row, 1, r + 1)
        cell(ws, row, 2, f"=INDEX({CNT}!$B${CNT_FIRST}:$F${CNT_FIRST + n_symbols - 1},MATCH({pref('scatter')},{cnt_rng},0),A{row})", S.LINK)
        cell(ws, row, 3, f"=INDEX({CNT}!$B${CNT_LEN_ROW}:$F${CNT_LEN_ROW},A{row})", S.LINK)
        cell(ws, row, 4, f"={pref('rows')}*B{row}/C{row}", S.LINK, S.PROB)

    ws.cell(row=11, column=1, value="Visible pattern enumeration (1 = visible)").font = S.SUBTITLE
    header(ws, 12, 1, ["Pattern"] + [f"R{i + 1}" for i in range(n_reels)] + ["Count", "Probability"])
    for mask in range(1 << n_reels):
        row = SCAT_ENUM_FIRST + mask
        cell(ws, row, 1, mask)
        bits = [(mask >> i) & 1 for i in range(n_reels)]
        for i, b in enumerate(bits):
            cell(ws, row, 2 + i, b)
        cell(ws, row, 7, f"=SUM(B{row}:F{row})")
        prob = "*".join(f"IF({'BCDEF'[i]}{row}=1,$D${5 + i},1-$D${5 + i})" for i in range(n_reels))
        cell(ws, row, 8, f"={prob}", fmt=S.PROB)
    enum_last = SCAT_ENUM_FIRST + (1 << n_reels) - 1

    SCAT_DIST_FIRST = enum_last + 4
    ws.cell(row=SCAT_DIST_FIRST - 2, column=1, value="SCATTER count distribution and pays").font = S.SUBTITLE
    header(ws, SCAT_DIST_FIRST - 1, 1, ["Count", "Probability", "Pay (x total bet)", "Contribution", "1 in N"])
    for k in range(n_reels + 1):
        row = SCAT_DIST_FIRST + k
        cell(ws, row, 1, k)
        cell(ws, row, 2, f"=SUMIF($G${SCAT_ENUM_FIRST}:$G${enum_last},A{row},$H${SCAT_ENUM_FIRST}:$H${enum_last})", fmt=S.PROB)
        cell(ws, row, 3, f"=IF(A{row}=0,0,INDEX({PAY}!$B${SCAT_PAY_ROW}:$F${SCAT_PAY_ROW},A{row}))", S.LINK, S.NUM)
        cell(ws, row, 4, f"=B{row}*C{row}", fmt=S.PROB)
        cell(ws, row, 5, f"=IF(B{row}>0,1/B{row},\"-\")", fmt=S.NUM2)
    SCAT_TOTAL_ROW = SCAT_DIST_FIRST + n_reels + 1
    t = SCAT_TOTAL_ROW
    cell(ws, t, 1, "Total", S.BOLD, fill=S.TOTAL_FILL)
    cell(ws, t, 2, f"=SUM(B{SCAT_DIST_FIRST}:B{t - 1})", S.BOLD, S.PROB, S.TOTAL_FILL)
    cell(ws, t, 3, None, fill=S.TOTAL_FILL)
    cell(ws, t, 4, f"=SUM(D{SCAT_DIST_FIRST}:D{t - 1})", S.BOLD, S.PCT4, S.TOTAL_FILL)
    cell(ws, t, 5, "← scatter RTP", S.NOTE, fill=S.TOTAL_FILL)

    SCAT_TRIG_ROW = t + 2
    cell(ws, SCAT_TRIG_ROW, 1, "FS trigger probability", S.BOLD)
    cell(ws, SCAT_TRIG_ROW, 2,
         f'=SUMIF(A{SCAT_DIST_FIRST}:A{t - 1},">="&{pref("trigger")},B{SCAT_DIST_FIRST}:B{t - 1})', S.LINK, S.PROB)
    cell(ws, SCAT_TRIG_ROW + 1, 1, "FS trigger 1 in N", S.BOLD)
    cell(ws, SCAT_TRIG_ROW + 1, 2, f"=1/B{SCAT_TRIG_ROW}", fmt=S.NUM2)


def fs_calc(wb, n_normal):
    ws = wb.create_sheet("Free Spin Calc")
    title(ws, "Free Spin Calc (exact)",
          "Free spins use the same reels as the base game, with every win × multiplier. Each FS spin retriggers with probability p and adds R more spins → "
          "expected total spins = N / (1 − R·p) (branching process, finite when R·p < 1).")
    widths(ws, {"A": 42, "B": 18, "C": 60})
    header(ws, 3, 1, ["Item", "Value", "Formula / notes"])
    rows = [
        ("FS trigger probability p", f"={SCAT}!B{SCAT_TRIG_ROW}", S.PROB, "Scatter Calc sheet"),
        ("Spins awarded N", f"={pref('award')}", S.NUM, "Base Parameters"),
        ("Retrigger spins R", f"={pref('retrigger')}", S.NUM, "Base Parameters"),
        ("Multiplier M", f"={pref('mult')}", S.NUM, "Base Parameters"),
        ("Branching rate R·p", "=B6*B4", "0.000000", "Must be below 1"),
        ("Expected total FS spins", '=IF(B8<1,B5/(1-B8),"Diverges")', "0.0000", "N / (1 − R·p)"),
        ("Base RTP (line + scatter)", f"={LINE}!AA{LINE_TOTAL_ROW}+{SCAT}!D{SCAT_TOTAL_ROW}", S.PCT4, "Expected pay per spin (x total bet)"),
        ("Expected pay per FS spin (x total bet)", "=B10*B7", "0.000000", "Base RTP × M"),
        ("Expected total win per FS feature (x total bet)", "=B9*B11", "0.00", "Expected total spins × pay per spin"),
        ("FS RTP contribution", "=B4*B12", S.PCT4, "p × expected feature win"),
    ]
    for i, (k, f, fmt, note) in enumerate(rows, start=4):
        cell(ws, i, 1, k, S.BOLD)
        cell(ws, i, 2, f, S.LINK if "!" in f else S.BODY, fmt)
        cell(ws, i, 3, note)


def rtp_breakdown(wb, n_normal):
    ws = wb.create_sheet("RTP Breakdown")
    title(ws, "RTP Breakdown")
    widths(ws, {"A": 30, "B": 16, "C": 16, "D": 40})
    header(ws, 4, 1, ["Component", "RTP", "Share", "Source"])
    rows = [
        ("Line pays", f"={LINE}!AA{LINE_TOTAL_ROW}", "Line Calc"),
        ("Scatter pays", f"={SCAT}!D{SCAT_TOTAL_ROW}", "Scatter Calc"),
        ("Base game subtotal", "=B5+B6", ""),
        ("Free spins", f"={FS}!B13", "Free Spin Calc"),
        ("Total RTP", "=B7+B8", ""),
    ]
    for i, (k, f, src) in enumerate(rows, start=5):
        is_total = k in ("Base game subtotal", "Total RTP")
        fill = S.TOTAL_FILL if is_total else None
        cell(ws, i, 1, k, S.BOLD, fill=fill)
        cell(ws, i, 2, f, S.LINK if "!" in f else S.BOLD, S.PCT4, fill)
        cell(ws, i, 3, f"=B{i}/$B$9", fmt=S.PCT2, fill=fill)
        cell(ws, i, 4, src, fill=fill)
    cell(ws, 11, 1, "Target RTP", S.BOLD)
    cell(ws, 11, 2, f"={pref('target')}", S.LINK, S.PCT2)
    cell(ws, 12, 1, "Difference from target", S.BOLD)
    cell(ws, 12, 2, "=B9-B11", fmt=S.PCT4)
    cell(ws, 13, 1, "Verdict (within ±0.1%)", S.BOLD)
    cell(ws, 13, 2, '=IF(ABS(B12)<=0.001,"OK","Needs tuning")')


# ---------------------------------------------------------------- simulation / verification
def sim_sheet(wb, sim):
    ws = wb.create_sheet("Simulation Results")
    mc, fc, th = sim["monteCarlo"], sim.get("fullCycle"), sim["theory"]
    title(ws, "Simulation Results (SDK slotsim)",
          f"Source: slotsim report / math {sim['model']['id']} v{sim['model']['version']} / generated {sim['generatedAt']}. "
          "Values are pasted from the SDK output (they don't recalculate). After changing the math, rerun make_specs.sh.")
    widths(ws, {"A": 34, "B": 18, "C": 4, "D": 34, "E": 18, "F": 4, "G": 16, "H": 16, "I": 16, "J": 16})

    header(ws, 4, 1, ["Monte Carlo", "Value"])
    mc_rows = [
        ("Rounds", mc["rounds"], S.NUM),                       # 5
        ("Total RTP", mc["rtp"], S.PCT4),                         # 6
        ("RTP 95% CI (±)", mc["rtp95HalfWidth"], S.PCT4),      # 7
        ("Base RTP", mc["baseRtp"], S.PCT4),                    # 8
        ("FS RTP", mc["freeSpinRtp"], S.PCT4),                    # 9
        ("Hit frequency", mc["hitFrequency"], S.PCT4),            # 10
        ("FS trigger frequency", mc["triggerFrequency"], S.PROB),     # 11
        ("Average FS spins", mc["avgFreeSpinsPerTrigger"], "0.000"),  # 12
        ("Average FS feature win (x total bet)", mc["avgFeatureWinMultiple"], S.NUM2),  # 13
        ("Standard deviation (per round)", mc["stdDev"], "0.000"),           # 14
        ("Max round win (x total bet)", mc["maxRoundWinMultiple"], S.NUM2),  # 15
        ("Seed", mc["seed"], None),                          # 16
    ]
    for i, (k, v, fmt) in enumerate(mc_rows, start=5):
        cell(ws, i, 1, k, S.BOLD)
        cell(ws, i, 2, v, S.NOTE, fmt)

    header(ws, 4, 4, ["Full cycle (base game)", "Value"])
    fc_rows = []
    if fc:
        fc_rows = [
            ("Combinations", fc["combinations"], S.NUM),        # 5
            ("Base RTP (exact)", fc["baseRtp"], "0.0000000000%"),   # 6
            ("Hit frequency (exact)", fc["hitFrequency"], S.PCT4),     # 7
            ("Max base win (x total bet)", fc["maxWinMultiple"], S.NUM2),  # 8
            ("FS trigger probability (exact)", sum(fc["scatterHistogram"][3:]) / fc["combinations"], "0.0000000000"),  # 9
        ]
    for i, (k, v, fmt) in enumerate(fc_rows, start=5):
        cell(ws, i, 4, k, S.BOLD)
        cell(ws, i, 5, v, S.NOTE, fmt)

    header(ws, 11, 4, ["SDK theory (TheoreticalCalculator)", "Value"])
    for i, (k, key) in enumerate([("Total RTP", "totalRtp"), ("Base RTP", "baseRtp"), ("FS trigger probability", "triggerProbability")], start=12):
        cell(ws, i, 4, k, S.BOLD)
        cell(ws, i, 5, th[key], S.NOTE, "0.0000000000")

    r0 = 19
    ws.cell(row=r0 - 1, column=1, value="Win distribution (round win ÷ total bet)").font = S.SUBTITLE
    header(ws, r0, 1, ["Range", "MC probability", "", "MC 1 in N", "Full cycle probability (base only)"])
    for i, b in enumerate(mc["buckets"]):
        row = r0 + 1 + i
        cell(ws, row, 1, b["label"])
        cell(ws, row, 2, b["probability"], S.NOTE, S.PCT4)
        cell(ws, row, 4, f'=IF(B{row}>0,1/B{row},"-")', fmt=S.NUM2)
        if fc:
            cell(ws, row, 5, fc["buckets"][i]["probability"], S.NOTE, S.PCT4)

    chart = BarChart()
    chart.title = "Win distribution (Monte Carlo)"
    chart.y_axis.title = "Probability"
    chart.y_axis.numFmt = "0%"
    chart.x_axis.title = "Round win / total bet"
    data = Reference(ws, min_col=2, min_row=r0, max_row=r0 + len(mc["buckets"]))
    cats = Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + len(mc["buckets"]))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.height = 8
    chart.width = 18
    ws.add_chart(chart, "G4")


def verify(wb, sim):
    ws = wb.create_sheet("Verification")
    title(ws, "Verification (Excel formulas vs SDK)",
          "Checks that the theoretical values from the Excel formulas agree with the SDK's independent calculations (evaluator enumeration and Monte Carlo).")
    widths(ws, {"A": 34, "B": 18, "C": 18, "D": 16, "E": 12, "F": 50})
    header(ws, 4, 1, ["Check", "Excel theory", "SDK", "Difference / z", "Verdict", "Criterion"])
    rows = [
        ("Base RTP vs full cycle", f"={RTP}!B7", f"={SIMS}!E6", "=B5-C5", '=IF(ABS(D5)<1E-9,"PASS","FAIL")',
         "Exact match (|difference| < 1e-9). Formula vs full enumeration of every stop combination by the actual evaluator"),
        ("FS trigger vs full cycle", f"={SCAT}!B{SCAT_TRIG_ROW}", f"={SIMS}!E9", "=B6-C6", '=IF(ABS(D6)<1E-12,"PASS","FAIL")',
         "Exact match"),
        ("Total RTP vs SDK theory", f"={RTP}!B9", f"={SIMS}!E12", "=B7-C7", '=IF(ABS(D7)<1E-9,"PASS","FAIL")',
         "Excel formulas vs C# TheoreticalCalculator"),
        ("Total RTP vs Monte Carlo", f"={RTP}!B9", f"={SIMS}!B6", f"=(C8-B8)/({SIMS}!B14/SQRT({SIMS}!B5))",
         '=IF(ABS(D8)<4,"PASS","FAIL")', "Within ±4σ in standard error (z value)"),
    ]
    fmts = ["0.0000000000%", "0.0000000000", "0.0000000000%", "0.0000%"]
    for i, (k, a, b, d, v, crit) in enumerate(rows, start=5):
        cell(ws, i, 1, k, S.BOLD)
        cell(ws, i, 2, a, S.LINK, fmts[i - 5])
        cell(ws, i, 3, b, S.LINK, fmts[i - 5])
        cell(ws, i, 4, d, fmt="0.00E+00" if i < 8 else "0.00")
        cell(ws, i, 5, v, S.BOLD)
        cell(ws, i, 6, crit)
    green = PatternFill("solid", fgColor="C8E6C9")
    red = PatternFill("solid", fgColor="FFCDD2")
    ws.conditional_formatting.add("E5:E8", CellIsRule(operator="equal", formula=['"PASS"'], fill=green))
    ws.conditional_formatting.add("E5:E8", CellIsRule(operator="equal", formula=['"FAIL"'], fill=red))
    ws["A10"] = "If you change the math in Excel, run export_math.py → make_specs.sh so the SDK-side values are recomputed."
    ws["A10"].font = S.NOTE
    ws["E5"].comment = Comment("Theory and the actual evaluator's exhaustive enumeration must match to machine precision.", "TakuroidX")


if __name__ == "__main__":
    main()
